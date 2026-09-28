import openpyxl
import json

wb = openpyxl.load_workbook(r"c:\Users\jasra\OneDrive\Desktop\Qoutely\qoutely\UPdated Calculator-07.02.2024.xlsx", data_only=True)
s2 = wb["Sheet2"]

records = []
for r in range(4, s2.max_row + 1):
    code = s2.cell(r, 2).value
    section = s2.cell(r, 3).value
    desc = s2.cell(r, 4).value
    category = s2.cell(r, 5).value
    flexa = s2.cell(r, 6).value
    stfi = s2.cell(r, 8).value
    eq = s2.cell(r, 10).value
    terrorism = s2.cell(r, 12).value
    
    if code is not None and desc is not None:
        records.append({
            "code": str(code).strip(),
            "section": str(section).strip() if section else "",
            "description": str(desc).strip(),
            "category": int(category) if category is not None and str(category).isdigit() else 1,
            "flexa_rate": float(flexa) if flexa is not None else 0.0,
            "stfi_rate": float(stfi) if stfi is not None else 0.15,
            "eq_rate": float(eq) if eq is not None else 0.1,
            "terrorism_rate": float(terrorism) if terrorism is not None else 0.15,
        })

print(f"Extracted {len(records)} occupancies from Calculator Sheet2.")
print("Sample record:", records[0])
print("Sample record (4002):", [r for r in records if r["code"] == "4002"])

# Check IIB Loss Cost - Schedule 3
wb_iib = openpyxl.load_workbook(r"c:\Users\jasra\OneDrive\Desktop\Qoutely\qoutely\IIB Loss Cost - Schedule 3.xlsx", data_only=True)
s_iib = wb_iib["Sheet2"]
iib_records = []
for r in range(10, s_iib.max_row + 1):
    c = s_iib.cell(r, 2).value
    d = s_iib.cell(r, 3).value
    lc = s_iib.cell(r, 4).value
    if c is not None and d is not None:
        iib_records.append({
            "code": str(c).strip(),
            "description": str(d).strip(),
            "loss_cost": float(lc) if lc is not None else 0.0
        })
print(f"Extracted {len(iib_records)} records from IIB Schedule 3.")
