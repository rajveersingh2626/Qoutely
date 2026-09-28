import openpyxl
import json
import os

os.makedirs(r"c:\Users\jasra\OneDrive\Desktop\Qoutely\src\data", exist_ok=True)

# Load occupancies from Updated Calculator Sheet2
wb = openpyxl.load_workbook(r"c:\Users\jasra\OneDrive\Desktop\Qoutely\qoutely\UPdated Calculator-07.02.2024.xlsx", data_only=True)
s2 = wb["Sheet2"]

# Load IIB Loss costs from Schedule 3
wb_iib = openpyxl.load_workbook(r"c:\Users\jasra\OneDrive\Desktop\Qoutely\qoutely\IIB Loss Cost - Schedule 3.xlsx", data_only=True)
s_iib = wb_iib["Sheet2"]
loss_costs = {}
for r in range(10, s_iib.max_row + 1):
    c = s_iib.cell(r, 2).value
    lc = s_iib.cell(r, 4).value
    if c is not None and lc is not None:
        loss_costs[str(c).strip()] = float(lc)

occupancies = []
for r in range(4, s2.max_row + 1):
    code = s2.cell(r, 2).value
    section = s2.cell(r, 3).value
    desc = s2.cell(r, 4).value
    cat = s2.cell(r, 5).value
    flexa = s2.cell(r, 6).value
    stfi = s2.cell(r, 8).value
    eq = s2.cell(r, 10).value
    terr = s2.cell(r, 12).value
    
    if code is not None and desc is not None:
        code_str = str(code).strip()
        desc_str = str(desc).strip()
        section_str = str(section).strip() if section else "IV"
        try:
            category_num = int(cat) if cat is not None else 1
        except:
            category_num = 1
            
        # Determine broad category tag
        tag = "Manufacturing"
        d_lower = desc_str.lower()
        if any(w in d_lower for w in ["godown", "silo", "storage", "warehouse", "depot"]):
            tag = "Warehouse"
        elif any(w in d_lower for w in ["shop", "showroom", "retail", "store", "mall", "market", "bazaar"]):
            tag = "Retail"
        elif any(w in d_lower for w in ["office", "bank", "meeting", "administrative"]):
            tag = "Office"
        elif any(w in d_lower for w in ["hotel", "restaurant", "cafe", "resort", "club", "mess", "boarding"]):
            tag = "Hospitality"
        elif any(w in d_lower for w in ["hospital", "clinic", "diagnostic", "nursing"]):
            tag = "Healthcare"
        elif any(w in d_lower for w in ["school", "college", "university", "institute", "educational"]):
            tag = "Education"
        elif any(w in d_lower for w in ["hazardous", "fireworks", "explosive", "ammunition", "celluloid", "lpg", "petroleum", "solvent"]):
            tag = "Hazardous"
            
        # Extract keywords
        raw_words = [w.strip(",.()[]{}/*&") for w in desc_str.split() if len(w) > 3]
        keywords = list(set([w.lower() for w in raw_words if w.lower() not in ["with", "from", "other", "using", "under", "where", "into", "their"]]))[:8]

        occupancies.append({
            "code": code_str,
            "section": section_str,
            "description": desc_str,
            "category": category_num,
            "category_tag": tag,
            "flexa_rate": round(float(flexa), 4) if flexa is not None else 0.5,
            "stfi_rate": round(float(stfi), 4) if stfi is not None else 0.15,
            "eq_rate": round(float(eq), 4) if eq is not None else 0.10,
            "terrorism_rate": round(float(terr), 4) if terr is not None else 0.15,
            "loss_cost": loss_costs.get(code_str, round(float(flexa) * 0.85, 4) if flexa else 0.35),
            "keywords": keywords,
            "source_doc": "IIB Loss Cost Schedule 3 / AIFT 2001 Section " + section_str
        })

with open(r"c:\Users\jasra\OneDrive\Desktop\Qoutely\src\data\occupancies.json", "w", encoding="utf-8") as f:
    json.dump(occupancies, f, indent=2, ensure_ascii=False)

print(f"Saved {len(occupancies)} occupancies to src/data/occupancies.json")
